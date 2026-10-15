"""A generated library and clients for the MCP tests: the SDK's own stdio and
streamable HTTP clients against real `docgrep mcp` and `docgrep serve`
processes with the deterministic backend."""

import json
import os
import socket
import subprocess
import sys
import time
from collections.abc import AsyncIterator, Awaitable, Callable
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any

import anyio
import httpx
from docx.document import Document
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from mcp.client.streamable_http import streamable_http_client
from mcp.types import CallToolResult, TextContent

from docgrep.api.container import create_container
from docgrep.settings import Settings
from docx_support import docx_bytes, write_docx
from pdf_support import pdf_bytes, prose, write_pdf
from resume_support import close

QUERY = "violet kraken lantern"
DUPLICATE = "The copper heron sleeps under the old stone bridge at dawn."
CORRUPT = bytes(range(256)) * 8


def _memo(document: Document) -> None:
    document.add_heading("Expedition", 1)
    document.add_paragraph("The harbour was calm before the voyage.")
    document.add_heading("Storm", 2)
    document.add_paragraph("The violet kraken lantern rose from the deep sea.")


def write_library(corpus: Path) -> None:
    """A paginated PDF, a DOCX (no pages), two copies of one note and a
    corrupt EPUB that fails to parse."""
    write_pdf(corpus / "paper.pdf", pdf_bytes([prose(1), prose(2)]))
    write_docx(corpus / "memo.docx", docx_bytes(_memo, title="Sea Memo"))
    (corpus / "notes").mkdir(parents=True, exist_ok=True)
    (corpus / "notes" / "a.txt").write_text(DUPLICATE, encoding="utf-8")
    (corpus / "notes" / "b.txt").write_text(DUPLICATE, encoding="utf-8")
    (corpus / "broken.epub").write_bytes(CORRUPT)


def settings_for(root: Path) -> Settings:
    return Settings(
        corpus_root=root / "corpus",
        data_dir=root / "data",
        backend="fake",
        chunk_size=64,
        chunk_overlap=8,
    )


def environment(root: Path, **extra: str) -> dict[str, str]:
    """The environment of a docgrep process for this library, and nothing of
    the caller's DOCGREP_ settings."""
    env = {k: v for k, v in os.environ.items() if not k.startswith("DOCGREP_")}
    env.update(
        DOCGREP_CORPUS_ROOT=str(root / "corpus"),
        DOCGREP_DATA_DIR=str(root / "data"),
        DOCGREP_BACKEND="fake",
        DOCGREP_CHUNK_SIZE="64",
        DOCGREP_CHUNK_OVERLAP="8",
        HF_HUB_OFFLINE="1",
    )
    env.update(extra)
    return env


def built_library(root: Path) -> Settings:
    """Write the library and build its index in this process, then release it."""
    write_library(root / "corpus")
    settings = settings_for(root)
    services = create_container(settings)
    summary = services.build.run()
    assert (summary.indexed, summary.failed) == (4, 1), summary
    close(services)
    return settings


def run[T](work: Callable[[], Awaitable[T]]) -> T:
    return anyio.run(work)


@asynccontextmanager
async def stdio_session(root: Path, **extra: str) -> AsyncIterator[ClientSession]:
    """A session with `docgrep mcp`; its standard error goes to mcp-stderr.log."""
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "docgrep.api.cli", "mcp"],
        env=environment(root, **extra),
        cwd=str(root),
    )
    with (root / "mcp-stderr.log").open("a", encoding="utf-8") as errlog:
        async with stdio_client(params, errlog=errlog) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                yield session


@asynccontextmanager
async def http_session(url: str) -> AsyncIterator[ClientSession]:
    async with streamable_http_client(url) as (read, write, _session_id):
        async with ClientSession(read, write) as session:
            await session.initialize()
            yield session


def free_port() -> int:
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port: int = probe.getsockname()[1]
        return port


class Serve:
    """A `docgrep serve` process on a free loopback port; stopped on exit."""

    def __init__(self, root: Path, **extra: str) -> None:
        self.port = free_port()
        self.base = f"http://127.0.0.1:{self.port}"
        self.log = (root / f"serve-{self.port}.log").open("w", encoding="utf-8")
        self.process = subprocess.Popen(
            [sys.executable, "-m", "docgrep.api.cli", "serve"],
            env=environment(
                root, DOCGREP_HOST="127.0.0.1", DOCGREP_PORT=str(self.port), **extra
            ),
            cwd=str(root),
            stdout=self.log,
            stderr=self.log,
        )

    def __enter__(self) -> "Serve":
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            try:
                if httpx.get(f"{self.base}/health", timeout=1).status_code == 200:
                    return self
            except httpx.HTTPError:
                pass
            if self.process.poll() is not None:
                break
            time.sleep(0.2)
        self.__exit__()
        raise RuntimeError("docgrep serve did not start")

    def __exit__(self, *_: object) -> None:
        self.process.terminate()
        try:
            self.process.wait(timeout=20)
        except subprocess.TimeoutExpired:
            self.process.kill()
            self.process.wait()
        self.log.close()


def structured(result: CallToolResult) -> dict[str, Any]:
    assert result.structuredContent is not None
    return dict(result.structuredContent)


def text(result: CallToolResult) -> str:
    (block,) = result.content
    assert isinstance(block, TextContent)
    return block.text


def without_timing(body: dict[str, Any]) -> dict[str, Any]:
    """A search body without took_ms, the only field that differs per call."""
    return {key: value for key, value in body.items() if key != "took_ms"}


def snapshot(*roots: Path) -> dict[str, tuple[int, int]]:
    return {
        str(path): (path.stat().st_size, path.stat().st_mtime_ns)
        for root in roots
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def status_json(root: Path) -> dict[str, Any]:
    """`docgrep status --json` for this library, in a separate process."""
    done = subprocess.run(
        [sys.executable, "-m", "docgrep.api.cli", "status", "--json"],
        env=environment(root),
        capture_output=True,
        text=True,
        check=True,
        timeout=60,
    )
    report: dict[str, Any] = json.loads(done.stdout)
    return report
