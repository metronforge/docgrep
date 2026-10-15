"""DOCX sources through builds and updates: a library that indexed them as
unsupported picks them up on update, and a missing python-docx only fails
them until the next update (UC-3 step 6.3, UC-4)."""

import sys
from pathlib import Path
from typing import Any

import pytest
from docx.document import Document

from docgrep.adapters.manifest import SQLiteManifest
from docgrep.domain.models import FileType
from docx_support import docx_bytes, write_docx
from resume_support import close, dump_points, settings_for
from update_support import update_services


def _report(document: Document) -> None:
    document.add_heading("Findings", 1)
    document.add_paragraph("marker7 the copper heron sleeps under the bridge")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "marker8 cell"
    table.cell(0, 1).text = "value"


def _corpus(root: Path) -> None:
    root.mkdir(parents=True)
    (root / "note.txt").write_text("marker1 plain words here", encoding="utf-8")
    write_docx(root / "report.docx", docx_bytes(_report, title="Field Report"))
    (root / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0legacy word")


def _built_without_docx(tmp_path: Path) -> Any:
    _corpus(tmp_path / "corpus")
    settings = settings_for(tmp_path / "corpus", tmp_path / "data")
    services = update_services(settings)
    build: Any = services.build
    build.parsers = {t: p for t, p in build.parsers.items() if t is not FileType.DOCX}
    assert build.run().failed == 0
    close(services)
    return settings


def _update(settings: Any) -> Any:
    services = update_services(settings)
    try:
        return services.update.run()
    finally:
        close(services)


def test_docx_skipped_as_unsupported_is_picked_up_by_update(tmp_path: Path) -> None:
    settings = _built_without_docx(tmp_path)
    manifest = SQLiteManifest(settings.data_dir)
    for path in ("report.docx", "legacy.doc"):
        record = manifest.get_source(path)
        assert record is not None and record.status == "skipped:unsupported"

    summary = _update(settings)

    plan = summary.plan
    assert (plan.newly_supported, plan.modified) == (1, 1)
    assert (plan.content_changed, plan.parser_changed, plan.retried) == (0, 0, 0)
    report = manifest.get_source("report.docx")
    assert report is not None and report.status == "indexed"
    assert report.source.file_type is FileType.DOCX
    assert report.source.parser_version == "1"
    assert report.source.title == "Field Report"
    legacy = manifest.get_source("legacy.doc")
    assert legacy is not None and legacy.status == "skipped:unsupported"
    points = [p[1] for p in dump_points(settings) if p[1]["path"] == "report.docx"]
    assert points
    assert {p["page_number"] for p in points} == {None}
    assert {tuple(p["heading_path"]) for p in points} == {("Findings",)}
    assert _update(settings).no_changes


def test_a_missing_library_fails_docx_only_until_the_next_update(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _corpus(tmp_path / "corpus")
    settings = settings_for(tmp_path / "corpus", tmp_path / "data")
    with monkeypatch.context() as patched:
        patched.setitem(sys.modules, "docx", None)
        services = update_services(settings)
        summary = services.build.run()
        close(services)

    assert (summary.indexed, summary.failed) == (1, 1)
    manifest = SQLiteManifest(settings.data_dir)
    report = manifest.get_source("report.docx")
    assert report is not None and report.status == "failed:parse_error"
    assert "python-docx is not installed" in (report.error or "")

    retried = _update(settings)

    assert retried.plan.retried == 1
    report = manifest.get_source("report.docx")
    assert report is not None and report.status == "indexed"
