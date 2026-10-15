"""DOCX documents through the composition root: build, search and manifest."""

from pathlib import Path

from docx.document import Document
from fastapi.testclient import TestClient

from docgrep.adapters.manifest import SQLiteManifest
from docgrep.api.http import create_app
from docgrep.settings import Settings
from docx_support import OLE_HEADER, docx_bytes, write_docx


def _memo(document: Document) -> None:
    document.add_heading("Expedition", 1)
    document.add_paragraph("The harbour was calm.")
    document.add_heading("Storm", 2)
    document.add_paragraph("The violet kraken lantern rose from the deep sea.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Crew"
    table.cell(0, 1).text = "twelve sailors"


def test_docx_documents_are_indexed_searched_and_rejected(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    write_docx(
        corpus / "memo.docx", docx_bytes(_memo, title="Sea Memo", author="Jane Writer")
    )
    write_docx(corpus / "locked.docx", OLE_HEADER)
    write_docx(corpus / "blank.docx", docx_bytes(lambda d: None))
    (corpus / "old.doc").write_bytes(OLE_HEADER)
    settings = Settings(
        corpus_root=corpus,
        data_dir=tmp_path / "data",
        backend="fake",
        chunk_size=64,
        chunk_overlap=8,
    )

    with TestClient(create_app(settings), base_url="http://127.0.0.1") as client:
        built = client.post("/index/build")
        assert built.status_code == 200, built.text
        assert {k: built.json()[k] for k in ("indexed", "skipped", "failed")} == {
            "indexed": 1,
            "skipped": 2,
            "failed": 1,
        }
        response = client.post(
            "/search", json={"query": "violet kraken lantern", "top_n": 1}
        )
        assert response.status_code == 200, response.text
        top = response.json()["results"][0]

    assert top["source"]["path"] == "memo.docx"
    assert top["source"]["file_type"] == "docx"
    assert top["source"]["title"] == "Sea Memo"
    assert top["source"]["author"] == "Jane Writer"
    text = top["passage"]["text"]
    assert "The violet kraken lantern rose from the deep sea." in text
    assert top["passage"]["heading_path"] == ["Expedition", "Storm"]
    assert top["passage"]["page_number"] is None
    manifest = SQLiteManifest(settings.data_dir)
    locked = manifest.get_source("locked.docx")
    blank = manifest.get_source("blank.docx")
    old = manifest.get_source("old.doc")
    assert locked is not None and locked.status == "failed:parse_error"
    assert "password-protected" in (locked.error or "")
    assert blank is not None and blank.status == "skipped:requires_ocr"
    assert old is not None and old.status == "skipped:unsupported"
