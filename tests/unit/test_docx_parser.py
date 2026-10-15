"""The DOCX parser on documents generated in the test (UC-03)."""

import io
import sys
from pathlib import Path

import pytest
from docx.document import Document
from docx.oxml.ns import qn

from docgrep.adapters.chunking import TokenChunker, WhitespaceTokenizer
from docgrep.adapters.parsers import DocxParser
from docgrep.domain.errors import ParseError
from docgrep.domain.models import FileType, ParsedDocument, Source
from docgrep.domain.protocols import DocumentParser
from docx_support import (
    OLE_HEADER,
    add_heading_style,
    docx_bytes,
    set_outline_level,
    tiny_png,
    write_docx,
)


def source(path: str = "doc.docx", file_type: FileType = FileType.DOCX) -> Source:
    return Source(source_id=path, path=path, file_type=file_type)


def parse(root: Path, data: bytes, name: str = "doc.docx") -> ParsedDocument:
    write_docx(root / name, data)
    parser: DocumentParser = DocxParser(root)
    return parser.parse(source(name))


def failure(root: Path, data: bytes) -> ParseError:
    with pytest.raises(ParseError) as raised:
        parse(root, data)
    assert raised.value.category == "parse_error"
    assert raised.value.details["path"] == "doc.docx"
    return raised.value


def sections(document: ParsedDocument) -> list[tuple[tuple[str, ...], str]]:
    return [(s.heading_path, s.text) for s in document.sections]


def _headings(document: Document) -> None:
    document.add_paragraph("Before any heading.")
    document.add_heading("Report", 0)
    document.add_paragraph("Under the title.")
    document.add_heading("Part One", 1)
    document.add_paragraph("Opening of part one.")
    document.add_heading("Chapter 1", 2)
    document.add_paragraph("First chapter text.")
    document.add_paragraph("Second paragraph.")
    document.add_heading("   ", 2)
    document.add_paragraph("After an empty heading.")
    document.add_heading("Chapter 2", 2)
    document.add_paragraph("Second chapter text.")
    document.add_heading("Part Two", 1)
    document.add_paragraph("Item a", style="List Bullet")
    document.add_paragraph("Item b", style="List Number")


def test_headings_build_the_section_path_without_page_numbers(tmp_path: Path) -> None:
    document = parse(tmp_path, docx_bytes(_headings))

    assert sections(document) == [
        ((), "Before any heading."),
        (("Report",), "Under the title."),
        (("Report", "Part One"), "Opening of part one."),
        (
            ("Report", "Part One", "Chapter 1"),
            "First chapter text.\n\nSecond paragraph.\n\nAfter an empty heading.",
        ),
        (("Report", "Part One", "Chapter 2"), "Second chapter text."),
        (("Report", "Part Two"), "Item a\n\nItem b"),
    ]
    assert all(s.page_number is None for s in document.sections)
    assert not document.requires_ocr


def test_localized_heading_styles_are_found_by_outline_level(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        first = add_heading_style(document, "Заголовок 1", 1)
        second = add_heading_style(document, "Заголовок 2", 2)
        document.add_paragraph("Глава первая", style=first)
        document.add_paragraph("Текст главы.")
        document.add_paragraph("Раздел", style=second)
        document.add_paragraph("Текст раздела.")
        # Direct formatting: a paragraph made a level-1 heading by itself.
        direct = document.add_paragraph("Глава вторая")
        set_outline_level(direct._p, 0)
        document.add_paragraph("Ещё текст.")

    assert sections(parse(tmp_path, docx_bytes(build))) == [
        (("Глава первая",), "Текст главы."),
        (("Глава первая", "Раздел"), "Текст раздела."),
        (("Глава вторая",), "Ещё текст."),
    ]


def test_heading_style_ids_are_the_fallback_without_an_outline_level(
    tmp_path: Path,
) -> None:
    def build(document: Document) -> None:
        heading = document.styles["Heading 1"].element
        for pPr in heading.findall(qn("w:pPr")):
            for outline in pPr.findall(qn("w:outlineLvl")):
                pPr.remove(outline)
        document.add_heading("Named only", 1)
        document.add_paragraph("Body.")
        # Outline level 9 is body text: never a heading, whatever the style.
        body = document.add_heading("Not a heading", 2)
        set_outline_level(body._p, 9)

    assert sections(parse(tmp_path, docx_bytes(build))) == [
        (("Named only",), "Body.\n\nNot a heading"),
    ]


def test_paragraphs_and_tables_keep_document_order(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        document.add_heading("Prices", 1)
        document.add_paragraph("Before the table.")
        table = document.add_table(rows=3, cols=3)
        rows = [("Item", "Qty", "Price"), ("", "", ""), ("Tea", "2", "5")]
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
        document.add_paragraph("Between the tables.")
        second = document.add_table(rows=1, cols=2)
        second.cell(0, 0).text = "Total"
        second.cell(0, 1).text = "10"
        document.add_paragraph("After the tables.")

    assert sections(parse(tmp_path, docx_bytes(build))) == [
        (
            ("Prices",),
            "Before the table.\n\n"
            "Item | Qty | Price\nTea | 2 | 5\n\n"
            "Between the tables.\n\n"
            "Total | 10\n\n"
            "After the tables.",
        )
    ]


def test_merged_cells_are_read_once(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        table = document.add_table(rows=3, cols=3)
        wide = table.cell(0, 0).merge(table.cell(0, 2))
        wide.text = "Spans three columns"
        tall = table.cell(1, 0).merge(table.cell(2, 0))
        tall.text = "Spans two rows"
        table.cell(1, 1).text = "b1"
        table.cell(1, 2).text = "c1"
        table.cell(2, 1).text = "b2"
        table.cell(2, 2).text = "c2"

    (section,) = parse(tmp_path, docx_bytes(build)).sections
    assert section.text == "Spans three columns\nSpans two rows | b1 | c1\nb2 | c2"
    assert section.text.count("Spans") == 2


def test_nested_tables_are_flattened_in_order(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Outer A"
        holder = table.cell(0, 1)
        holder.text = "Outer B"
        inner = holder.add_table(rows=2, cols=2)
        inner.cell(0, 0).text = "Inner 1"
        inner.cell(0, 1).text = "Inner 2"
        inner.cell(1, 0).text = "Inner 3"
        inner.cell(1, 1).text = "Inner 4"
        table.cell(1, 0).text = "Outer C"
        table.cell(1, 1).text = "Outer D"

    (section,) = parse(tmp_path, docx_bytes(build)).sections
    assert section.text == (
        "Outer A | Outer B\nInner 1 | Inner 2\nInner 3 | Inner 4\nOuter C | Outer D"
    )


def test_out_of_scope_parts_are_ignored_without_failing(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        section = document.sections[0]
        section.header.add_paragraph("Running header")
        section.footer.add_paragraph("Page footer")
        paragraph = document.add_paragraph("Body with a comment.")
        document.add_comment(paragraph.runs, text="A reviewer comment")
        document.add_picture(io.BytesIO(tiny_png()))

    document = parse(tmp_path, docx_bytes(build))

    assert sections(document) == [((), "Body with a comment.")]


def test_tracked_deletions_are_left_out_and_insertions_kept(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        paragraph = document.add_paragraph("Kept ")
        p = paragraph._p
        p.append(_tracked("w:ins", "w:t", "inserted "))
        p.append(_tracked("w:del", "w:delText", "deleted "))
        p.append(_tracked(None, "w:t", "end."))

    (section,) = parse(tmp_path, docx_bytes(build)).sections
    assert section.text == "Kept inserted end."


def _tracked(wrapper: str | None, text_tag: str, text: str) -> object:
    from docx.oxml import OxmlElement

    run = OxmlElement("w:r")
    node = OxmlElement(text_tag)
    node.text = text
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(node)
    if wrapper is None:
        return run
    outer = OxmlElement(wrapper)
    outer.set(qn("w:id"), "1")
    outer.set(qn("w:author"), "Editor")
    outer.append(run)
    return outer


def test_core_properties_give_title_and_author(tmp_path: Path) -> None:
    data = docx_bytes(
        lambda d: d.add_paragraph("Body."), title="Annual Report", author="Jane Doe"
    )

    document = parse(tmp_path, data)

    assert (document.source.title, document.source.author) == (
        "Annual Report",
        "Jane Doe",
    )
    assert document.source.year is None


def test_an_empty_document_is_left_to_ocr(tmp_path: Path) -> None:
    document = parse(tmp_path, docx_bytes(lambda d: None))

    assert document.sections == () and document.requires_ocr


def test_a_document_with_only_an_image_is_left_to_ocr(tmp_path: Path) -> None:
    document = parse(
        tmp_path, docx_bytes(lambda d: d.add_picture(io.BytesIO(tiny_png())))
    )

    assert document.sections == () and document.requires_ocr


def test_a_truncated_zip_is_not_a_valid_zip(tmp_path: Path) -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Body."))

    error = failure(tmp_path, data[: len(data) // 2])

    assert error.details["reason"] == "not_a_zip"
    assert "not a valid ZIP archive" in str(error)


def test_a_text_file_named_docx_is_not_a_valid_zip(tmp_path: Path) -> None:
    error = failure(tmp_path, b"just some plain text, not a package")

    assert error.details["reason"] == "not_a_zip"


def test_a_zip_without_ooxml_content_types_is_rejected(tmp_path: Path) -> None:
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")

    assert failure(tmp_path, buffer.getvalue()).details["reason"] == "not_ooxml"


def test_an_ole_compound_file_is_reported_as_encrypted(tmp_path: Path) -> None:
    error = failure(tmp_path, OLE_HEADER)

    assert error.details["reason"] == "encrypted_content"
    assert "password-protected" in str(error)


def test_failure_messages_carry_no_absolute_path(tmp_path: Path) -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Body."))
    for broken in (data[:100], b"plain", OLE_HEADER):
        error = failure(tmp_path, broken)
        assert str(tmp_path) not in str(error)
        assert str(tmp_path) not in str(error.details)


def test_a_missing_library_is_a_recoverable_parse_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_docx(tmp_path / "doc.docx", docx_bytes(lambda d: d.add_paragraph("x")))
    # A None entry makes the import of python-docx fail, as if not installed.
    monkeypatch.setitem(sys.modules, "docx", None)

    with pytest.raises(ParseError) as raised:
        DocxParser(tmp_path).parse(source())

    assert raised.value.details["reason"] == "docx_library_missing"
    assert raised.value.details["path"] == "doc.docx"
    assert "hint" in raised.value.details


def test_a_source_of_another_type_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ParseError):
        DocxParser(tmp_path).parse(source("doc.doc", FileType.UNSUPPORTED))


def test_parsing_is_deterministic_and_chunks_keep_the_heading_path(
    tmp_path: Path,
) -> None:
    data = docx_bytes(_headings)
    first = parse(tmp_path, data)
    second = parse(tmp_path, data)
    chunker = TokenChunker(WhitespaceTokenizer(), chunk_size=8, chunk_overlap=2)

    assert first == second
    passages = chunker.chunk(first)
    assert passages and all(p.page_number is None for p in passages)
    assert passages[-1].heading_path == ("Report", "Part Two")


def test_parsing_writes_nothing_to_disk(tmp_path: Path) -> None:
    write_docx(tmp_path / "doc.docx", docx_bytes(_headings))
    before = sorted(tmp_path.rglob("*"))

    DocxParser(tmp_path).parse(source())

    assert sorted(tmp_path.rglob("*")) == before
